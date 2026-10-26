import os
import json
import re
from io import BytesIO
from datetime import datetime, timedelta
from typing import Any
from urllib.parse import quote

from dotenv import load_dotenv
from flask import Flask, flash, jsonify, redirect, render_template, request, session, url_for
from flask_cors import CORS
from google import genai
from google.genai import types
from supabase import create_client, Client
from werkzeug.utils import secure_filename
from docx import Document
from pypdf import PdfReader


load_dotenv()

app = Flask(__name__, template_folder='templates')
app.secret_key = 'ai_traveler_secret_key_2026'
CORS(app, resources={r"/api/*": {"origins": "*"}})

GEMINI_MODEL = os.getenv("GEMINI_MODEL", "gemini-3.5-flash-lite")
GEMINI_TIMEOUT_MS = int(os.getenv("GEMINI_TIMEOUT_MS", "60000"))

spots = [
    {'name': '雷門', 'emoji': '⛩️', 'tag': '景點', 'category': 'attraction', 'desc': '充滿歷史感的紅色大燈籠。', 'rating': 4.8},
    {'name': '明治神宮', 'emoji': '⛩️', 'tag': '景點', 'category': 'attraction', 'desc': '東京市中心的森林神社。', 'rating': 4.7},
    {'name': '一蘭拉麵', 'emoji': '🍜', 'tag': '美食', 'category': 'food', 'desc': '豚骨拉麵的經典代表作。', 'rating': 4.5},
    {'name': '築地市場', 'emoji': '🍣', 'tag': '美食', 'category': 'food', 'desc': '新鮮海鮮與在地小吃集中地。', 'rating': 4.7},
]

# ==========================================
# Supabase 設定
# ==========================================

SUPABASE_URL = "https://keiyqvrdiogknmmghzgq.supabase.co"
SUPABASE_KEY = "sb_publishable_tTq7TudmepTnN1oqQxEv1g__2pjRfFN"
supabase: Client = create_client(SUPABASE_URL, SUPABASE_KEY)


def get_gemini_client() -> genai.Client | None:
    api_key = os.getenv("GEMINI_API_KEY")
    if not api_key:
        return None
    return genai.Client(
        api_key=api_key,
        http_options=types.HttpOptions(timeout=GEMINI_TIMEOUT_MS),
    )


def parse_gemini_json(text: str) -> Any:
    cleaned = text.strip()
    fenced = re.search(r"```(?:json)?\s*(.*?)\s*```", cleaned, re.DOTALL)
    if fenced:
        cleaned = fenced.group(1)
    return json.loads(cleaned)


def generate_gemini_json(prompt: str, contents: Any = None) -> Any:
    client = get_gemini_client()
    if client is None:
        raise RuntimeError("GEMINI_API_KEY is not configured.")
    response = client.models.generate_content(
        model=GEMINI_MODEL,
        contents=contents if contents is not None else prompt,
        config={
            "system_instruction": (
                "你是旅遊規劃助手。只回傳有效 JSON，不要 Markdown 或額外文字。"
            ),
            "response_mime_type": "application/json",
            "temperature": 0.2,
            "max_output_tokens": 8192,
        },
    )
    if not response.text:
        raise ValueError("Gemini returned an empty response.")
    return parse_gemini_json(response.text)


def google_maps_search_url(place: str, city: str = '') -> str:
    query = ', '.join(part for part in (place, city) if part)
    return f"https://www.google.com/maps/search/?api=1&query={quote(query)}"


def google_maps_directions_url(places: list[str], city: str = '') -> str | None:
    destinations = [', '.join(part for part in (place, city) if part) for place in places if place]
    if not destinations:
        return None
    params = [
        'api=1',
        f'origin={quote(destinations[0])}',
        f'destination={quote(destinations[-1])}',
        'travelmode=transit',
    ]
    if len(destinations) > 2:
        params.append(f"waypoints={quote('|'.join(destinations[1:-1]))}")
    return f"https://www.google.com/maps/dir/?{'&'.join(params)}"


def normalize_activity(
    activity: Any,
    city: str,
    index: int,
    include_maps: bool = True,
) -> dict[str, str]:
    if isinstance(activity, dict):
        name = str(activity.get('name') or activity.get('place') or activity.get('location') or '').strip()
        place = str(activity.get('place') or activity.get('name') or name).strip()
        category = str(activity.get('category') or '活動').strip()
        normalized = {
            'time': str(activity.get('time') or f'時段 {index}').strip(),
            'name': name or place or '未命名行程',
            'place': place or name or '未指定地點',
            'address': str(activity.get('address') or '').strip(),
            'category': category,
            'description': str(activity.get('description') or '').strip(),
        }
    else:
        place = str(activity or '').strip()
        normalized = {
            'time': f'時段 {index}',
            'name': place or '未命名行程',
            'place': place or '未指定地點',
            'address': '',
            'category': '活動',
            'description': '',
        }
    if include_maps:
        normalized['mapsUrl'] = google_maps_search_url(
            normalized['address'] or normalized['place'], city
        )
    return normalized


def normalize_day(
    day_plan: Any,
    city: str,
    day_number: int,
    date_text: str,
    include_maps: bool = True,
) -> dict[str, Any]:
    source = day_plan if isinstance(day_plan, dict) else {}
    activities = source.get('activities') or []
    if not isinstance(activities, list):
        activities = [activities]
    normalized_activities = [
        normalize_activity(activity, city, index, include_maps=include_maps)
        for index, activity in enumerate(activities, start=1)
    ]
    places = [
        activity['address'] or activity['place']
        for activity in normalized_activities
        if activity['place'] != '未指定地點'
    ]
    normalized_day = {
        'day': day_number,
        'date': str(source.get('date') or date_text),
        'title': str(source.get('title') or f'{city} 第 {day_number} 天'),
        'description': str(source.get('description') or '').strip(),
        'activities': normalized_activities,
    }
    if include_maps:
        normalized_day['mapsUrl'] = google_maps_directions_url(places, city)
    return normalized_day


def get_authenticated_client() -> tuple[Any, Any]:
    user_id = session.get('user_id')
    access_token = session.get('access_token')
    refresh_token = session.get('refresh_token')
    if not user_id or not access_token or not refresh_token:
        return None, None

    try:
        user_client = create_client(SUPABASE_URL, SUPABASE_KEY)
        user_client.auth.set_session(access_token, refresh_token)
        return user_client, user_id
    except Exception:
        return None, None


def authentication_required() -> tuple[Any, Any, Any]:
    user_client, user_id = get_authenticated_client()
    if not user_client:
        return None, None, (jsonify({'error': '請先登入才能使用此功能。'}), 401)
    return user_client, user_id, None

# ==========================================
# 路由：首頁
# ==========================================
@app.route("/")
def home():
    # 從 session 中獲取當前用戶的資訊
    current_user = session.get('user')
    current_name = session.get('name')
    current_gender = session.get('gender')
    # 渲染首頁模板，並傳遞用戶資訊
    return render_template("index.html",
                           user=current_user,
                           name=current_name,
                           gender=current_gender)

# ==========================================
# 路由：其他頁面
# ==========================================
@app.route("/planner")
def planner():

    current_user = session.get('user')
    current_name = session.get('name')
    current_gender = session.get('gender')
    return render_template("planner.html",
                           user=current_user,
                           name=current_name,
                           gender=current_gender)

@app.route("/explore")
def explore():

    current_user = session.get('user')
    current_name = session.get('name')
    current_gender = session.get('gender')
    return render_template("explore.html",
                           user=current_user,
                           name=current_name,
                           gender=current_gender)

@app.route("/favorites")
def favorites():

    current_user = session.get('user')
    current_name = session.get('name')
    current_gender = session.get('gender')
    return render_template("favorites.html",
                           user=current_user,
                           name=current_name,
                           gender=current_gender,
                           favorites=[])

@app.route("/about")
def about():
    current_user = session.get('user')
    current_name = session.get('name')
    current_gender = session.get('gender')
    return render_template("about.html",
                           user=current_user,
                           name=current_name,
                           gender=current_gender)

@app.route("/recognize")
def recognize():
    current_user = session.get('user')
    current_name = session.get('name')
    current_gender = session.get('gender')
    return render_template("recognize.html",
                           user=current_user,
                           name=current_name,
                           gender=current_gender)


@app.route("/api/chat", methods=["POST"])
def chat():
    data = request.get_json(silent=True)
    if not isinstance(data, dict):
        return jsonify({"error": "請提供有效的 JSON 資料。"}), 400

    prompt = data.get("prompt")
    if not isinstance(prompt, str) or not prompt.strip():
        return jsonify({"error": "prompt 必須是不可為空的文字。"}), 400

    client = get_gemini_client()
    if client is None:
        app.logger.error("GEMINI_API_KEY is not configured.")
        return jsonify({"error": "Gemini API 尚未設定，請確認 GEMINI_API_KEY 環境變數。"}), 503

    try:
        response = client.models.generate_content(
            model=GEMINI_MODEL,
            contents=prompt.strip(),
        )
        return jsonify({"response": response.text})
    except Exception:
        app.logger.exception("Gemini API request failed.")
        return jsonify({"error": "Gemini API 呼叫失敗，請稍後再試。"}), 502


# ==========================================
# API：生成行程
# ==========================================
@app.route("/api/generate_itinerary", methods=['POST'])
def generate_itinerary():
    data = request.get_json(silent=True) or {}
    city = str(data.get('city', '')).strip()
    start_date_text = str(data.get('startDate', '')).strip()
    style = str(data.get('style', 'balanced')).strip() or 'balanced'
    budget = data.get('budget')
    try:
        days = max(1, min(14, int(data.get('days', 1))))
    except (TypeError, ValueError):
        return jsonify({'error': '旅行天數格式不正確。'}), 400
    if not city:
        return jsonify({'error': '請輸入目的地。'}), 400
    try:
        start_date = datetime.strptime(start_date_text, '%Y-%m-%d') if start_date_text else datetime.now()
    except ValueError:
        return jsonify({'error': '出發日期格式不正確。'}), 400

    prompt = f"""
規劃{city} {days}天旅行，日期從{start_date.strftime('%Y-%m-%d')}開始，風格：{style}。
只回傳 JSON，不要解釋。每天至少安排 4 個活動，最多安排 6 個活動，按照時間排序。
格式：{{"itinerary":[{{"day":1,"date":"YYYY-MM-DD","title":"短標題","description":"30字內摘要",
"activities":[{{"time":"09:00","name":"明確景點或店名","place":"明確地點名稱","address":"簡短地址",
"category":"景點","description":"20字內說明"}}]}}]}}
day 必須是 1 到 {days}，日期依序增加。category 只能是：食物、活動、購物、景點、酒店、交通、其他。
必須使用真實、明確的地點或店名，不要使用「探索市區」等泛稱。
"""
    try:
        result = generate_gemini_json(prompt)
        itinerary = result.get('itinerary') if isinstance(result, dict) else None
        if not isinstance(itinerary, list) or len(itinerary) != days:
            raise ValueError("Gemini returned an invalid itinerary shape.")
        if any(
            not isinstance(day_plan, dict)
            or not isinstance(day_plan.get('activities'), list)
            or len(day_plan['activities']) < 4
            for day_plan in itinerary
        ):
            raise ValueError("Gemini returned fewer than four activities for a day.")
        normalized_itinerary = [
            normalize_day(
                day_plan,
                city,
                day_number,
                (start_date + timedelta(days=day_number - 1)).strftime('%Y-%m-%d'),
                include_maps=False,
            )
            for day_number, day_plan in enumerate(itinerary, start=1)
        ]
        return jsonify({'itinerary': normalized_itinerary, 'source': 'gemini'})
    except RuntimeError:
        app.logger.error("GEMINI_API_KEY is not configured.")
        return jsonify({'error': 'Gemini API 尚未設定，請確認 GEMINI_API_KEY 環境變數。'}), 503
    except Exception:
        app.logger.exception("Gemini itinerary generation failed.")
        return jsonify({'error': 'AI 行程規劃失敗，請稍後再試。'}), 502

# ==========================================
# API：獲取探索景點
# ==========================================
@app.route("/api/explore_spots")
def get_explore_spots():
    category = request.args.get('category', 'all')
    category_name = {'all': '景點與美食', 'attraction': '景點', 'food': '美食'}.get(category)
    if category_name is None:
        return jsonify({'error': '無效的景點分類。'}), 400

    try:
        result = generate_gemini_json(
            f"""
請隨機推薦 6 個適合旅客探索的{category_name}。
每個推薦需包含真實且適合旅遊的名稱、emoji、tag、category、簡短 desc 與 1 到 5 的 rating。
category 只能是 attraction 或 food。
回傳 JSON：{{"spots":[{{"name":"名稱","emoji":"圖示","tag":"景點或美食","category":"attraction","desc":"介紹","rating":4.5}}]}}
"""
        )
        generated_spots = result.get('spots') if isinstance(result, dict) else None
        if not isinstance(generated_spots, list) or not generated_spots:
            raise ValueError("Gemini returned an invalid spots shape.")
        for spot in generated_spots:
            spot['mapsUrl'] = google_maps_search_url(str(spot.get('name', '')).strip())
        return jsonify({'spots': generated_spots, 'source': 'gemini'})
    except RuntimeError:
        filtered = spots if category == 'all' else [spot for spot in spots if spot['category'] == category]
        return jsonify({
            'spots': [{**spot, 'mapsUrl': google_maps_search_url(spot['name'])} for spot in filtered],
            'source': 'fallback',
        })
    except Exception:
        app.logger.exception("Gemini spot recommendation failed.")
        return jsonify({'error': 'AI 景點推薦失敗，請稍後再試。'}), 502


@app.get('/api/favorites')
def get_favorites():
    user_client, user_id, error = authentication_required()
    if error:
        return error
    try:
        favorite_rows = user_client.table('ai_recommendation_favorites').select(
            'id, recommendation_id, created_at'
        ).eq('user_id', user_id).order('created_at', desc=True).execute().data or []
        recommendation_ids = [row['recommendation_id'] for row in favorite_rows]
        if not recommendation_ids:
            return jsonify({'favorites': []})

        recommendations = user_client.table('ai_recommendations').select('*').in_(
            'id', recommendation_ids
        ).eq('user_id', user_id).execute().data or []
        recommendations_by_id = {item['id']: item for item in recommendations}
        favorites = []
        for favorite in favorite_rows:
            recommendation = recommendations_by_id.get(favorite['recommendation_id'])
            if recommendation:
                item = recommendation.get('recommendation_data') or {}
                item = dict(item) if isinstance(item, dict) else {}
                item.setdefault('name', recommendation.get('recommendation_text', '未命名推薦'))
                item['recommendationId'] = recommendation['id']
                item['favoriteId'] = favorite['id']
                favorites.append(item)
        return jsonify({'favorites': favorites})
    except Exception as error:
        return jsonify({'error': f'讀取收藏失敗：{error}'}), 500


@app.post('/api/upload_itinerary')
def upload_itinerary():
    uploaded_file = request.files.get('file')
    if uploaded_file is None or not uploaded_file.filename:
        return jsonify({'error': '請選擇要上傳的檔案。'}), 400

    filename = secure_filename(uploaded_file.filename)
    extension = os.path.splitext(filename)[1].lower()
    if extension not in {'.txt', '.doc', '.docx', '.pdf'}:
        return jsonify({'error': '只支援 TXT、Word（DOC/DOCX）與 PDF 檔案。'}), 415

    file_bytes = uploaded_file.read()
    if not file_bytes:
        return jsonify({'error': '檔案內容不可為空。'}), 400
    if len(file_bytes) > 10 * 1024 * 1024:
        return jsonify({'error': '檔案大小不可超過 10 MB。'}), 413

    try:
        if extension == '.txt':
            try:
                text = file_bytes.decode('utf-8-sig')
            except UnicodeDecodeError:
                text = file_bytes.decode('cp950')
        elif extension == '.docx':
            document = Document(BytesIO(file_bytes))
            paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
            table_rows = [
                ' | '.join(cell.text.strip() for cell in row.cells)
                for table in document.tables
                for row in table.rows
            ]
            text = '\n'.join(paragraphs + [row for row in table_rows if row])
        elif extension == '.pdf':
            reader = PdfReader(BytesIO(file_bytes))
            text = '\n'.join(page.extract_text() or '' for page in reader.pages).strip()
        else:
            return jsonify({'error': '舊式 DOC 檔案無法直接解析，請另存為 DOCX 後再上傳。'}), 415
    except Exception:
        app.logger.exception('Itinerary file parsing failed.')
        return jsonify({'error': '檔案解析失敗，請確認檔案格式正確。'}), 422

    text = text.strip()
    if not text:
        return jsonify({'error': '檔案中沒有可讀取的文字內容。'}), 422
    return jsonify({
        'fileName': filename,
        'fileType': extension[1:].upper(),
        'text': text[:100000],
    })


@app.post('/api/favorites')
def create_favorite():
    user_client, user_id, error = authentication_required()
    if error:
        return error
    data = request.get_json(silent=True) or {}
    name = str(data.get('name', '')).strip()
    if not name:
        return jsonify({'error': '缺少收藏內容。'}), 400

    try:
        recommendation = user_client.table('ai_recommendations').insert({
            'user_id': user_id,
            'recommendation_text': str(data.get('desc') or name),
            'recommendation_data': data,
            'source': str(data.get('source') or 'explore'),
        }).execute().data[0]
        favorite = user_client.table('ai_recommendation_favorites').insert({
            'user_id': user_id,
            'recommendation_id': recommendation['id'],
        }).execute().data[0]
        return jsonify({'favorite': {**data, 'recommendationId': recommendation['id'], 'favoriteId': favorite['id']}}), 201
    except Exception as error:
        return jsonify({'error': f'新增收藏失敗：{error}'}), 500


@app.delete('/api/favorites/<recommendation_id>')
def delete_favorite(recommendation_id):
    user_client, user_id, error = authentication_required()
    if error:
        return error
    try:
        user_client.table('ai_recommendation_favorites').delete().eq(
            'user_id', user_id
        ).eq('recommendation_id', recommendation_id).execute()
        return jsonify({'message': '已取消收藏。'})
    except Exception as error:
        return jsonify({'error': f'取消收藏失敗：{error}'}), 500


@app.post('/api/generations')
def create_generation_record():
    user_client, user_id, error = authentication_required()
    if error:
        return error
    data = request.get_json(silent=True) or {}
    itinerary = data.get('itinerary')
    city = str(data.get('city', '')).strip()
    if not city or not isinstance(itinerary, list) or not itinerary:
        return jsonify({'error': '缺少有效的行程資料。'}), 400

    try:
        record = user_client.table('ai_generation_records').insert({
            'user_id': user_id,
            'record_type': 'itinerary',
            'title': f'{city} 行程',
            'prompt': str(data.get('prompt') or f'規劃 {city} 的旅行行程'),
            'generated_content': json.dumps(itinerary, ensure_ascii=False),
            'context_data': {
                'city': city,
                'days': data.get('days'),
                'startDate': data.get('startDate'),
                'budget': data.get('budget'),
                'style': data.get('style'),
            },
            'model_name': 'demo-itinerary-generator',
            'request_payload': data,
            'response_payload': {'itinerary': itinerary},
        }).execute().data[0]
        return jsonify({'record': record}), 201
    except Exception as error:
        return jsonify({'error': f'儲存行程紀錄失敗：{error}'}), 500


@app.get('/api/generation-favorites')
def get_generation_favorites():
    user_client, user_id, error = authentication_required()
    if error:
        return error
    try:
        rows = user_client.table('ai_generation_favorites').select(
            'id, generation_record_id, created_at'
        ).eq('user_id', user_id).order('created_at', desc=True).execute().data or []
        record_ids = [row['generation_record_id'] for row in rows]
        records = []
        if record_ids:
            records = user_client.table('ai_generation_records').select('*').in_(
                'id', record_ids
            ).eq('user_id', user_id).execute().data or []
        records_by_id = {record['id']: record for record in records}
        return jsonify({
            'favorites': [
                {**row, 'record': records_by_id[row['generation_record_id']]}
                for row in rows if row['generation_record_id'] in records_by_id
            ]
        })
    except Exception as error:
        return jsonify({'error': f'讀取行程收藏失敗：{error}'}), 500


@app.post('/api/generation-favorites')
def create_generation_favorite():
    user_client, user_id, error = authentication_required()
    if error:
        return error
    data = request.get_json(silent=True) or {}
    generation_record_id = str(data.get('generationRecordId', '')).strip()
    if not generation_record_id:
        return jsonify({'error': '缺少生成紀錄 ID。'}), 400
    try:
        record_rows = user_client.table('ai_generation_records').select('id').eq(
            'id', generation_record_id
        ).eq('user_id', user_id).limit(1).execute().data or []
        if not record_rows:
            return jsonify({'error': '找不到可收藏的行程紀錄。'}), 404

        existing_rows = user_client.table('ai_generation_favorites').select('*').eq(
            'user_id', user_id
        ).eq('generation_record_id', generation_record_id).limit(1).execute().data or []
        if existing_rows:
            return jsonify({'favorite': existing_rows[0]}), 200

        favorite = user_client.table('ai_generation_favorites').insert({
            'user_id': user_id,
            'generation_record_id': generation_record_id,
        }).execute().data[0]
        return jsonify({'favorite': favorite}), 201
    except Exception as error:
        return jsonify({'error': f'收藏行程失敗：{error}'}), 500


@app.delete('/api/generation-favorites/<generation_record_id>')
def delete_generation_favorite(generation_record_id):
    user_client, user_id, error = authentication_required()
    if error:
        return error
    try:
        user_client.table('ai_generation_favorites').delete().eq(
            'user_id', user_id
        ).eq('generation_record_id', generation_record_id).execute()
        return jsonify({'message': '已取消行程收藏。'})
    except Exception as error:
        return jsonify({'error': f'取消行程收藏失敗：{error}'}), 500


@app.post('/api/search_history')
def create_search_history():
    user_client, user_id, error = authentication_required()
    if error:
        return error
    data = request.get_json(silent=True) or {}
    search_query = str(data.get('searchQuery', '')).strip()
    if not search_query:
        return jsonify({'error': '搜尋內容不可為空。'}), 400
    try:
        record = user_client.table('search_history').insert({
            'user_id': user_id,
            'search_query': search_query,
            'search_filters': data.get('filters') or {},
        }).execute().data[0]
        return jsonify({'record': record}), 201
    except Exception as error:
        return jsonify({'error': f'儲存搜尋紀錄失敗：{error}'}), 500


@app.post('/api/trip_insights')
def trip_insights():
    data = request.get_json(silent=True) or {}
    city = str(data.get('city', '')).strip()
    start_date_text = str(data.get('startDate', '')).strip()
    try:
        days = max(1, min(30, int(data.get('days', 1))))
        budget = max(0, float(data.get('budget') or 0))
    except (TypeError, ValueError):
        return jsonify({'error': '天數或預算格式不正確。'}), 400
    if not city:
        return jsonify({'error': '請先輸入目的地。'}), 400
    try:
        start_date = datetime.strptime(start_date_text, '%Y-%m-%d') if start_date_text else datetime.now()
    except ValueError:
        return jsonify({'error': '出發日期格式不正確。'}), 400

    hotel = round(budget * 0.35) if budget else days * 2200
    transport = round(budget * 0.12) if budget else days * 700
    food = round(budget * 0.25) if budget else days * 1000
    tickets = round(budget * 0.13) if budget else days * 500
    reserve = round(budget * 0.15) if budget else days * 600
    total = hotel + transport + food + tickets + reserve
    return jsonify({
        'weather': [
            {
                'day': day,
                'date': (start_date + timedelta(days=day - 1)).strftime('%Y-%m-%d'),
                'condition': '晴時多雲' if day % 2 else '短暫雨',
                'temperature': '18-24°C',
                'rain': 20 if day % 2 else 60,
            }
            for day in range(1, days + 1)
        ],
        'transport': {'summary': '建議使用大眾運輸，景點間平均約 25 分鐘。', 'estimatedCost': transport, 'tip': '尖峰時段請預留 15 分鐘轉乘緩衝。'},
        'accommodation': {'area': f'{city} 車站周邊', 'nightlyBudget': round(hotel / days), 'reason': '靠近交通節點，適合首次造訪與每日往返。'},
        'budget': {
            'hotel': hotel, 'transport': transport, 'food': food, 'tickets': tickets, 'reserve': reserve, 'total': total,
            'withinBudget': not budget or total <= budget,
            'message': '預算配置可行，已保留 15% 預備金。' if not budget or total <= budget else '目前配置可能超出預算，建議降低住宿或購物支出。',
        },
    })


@app.post('/api/replan')
def replan():
    data = request.get_json(silent=True) or {}
    city = str(data.get('city', '')).strip()
    reason = str(data.get('reason', '臨時變更')).strip() or '臨時變更'
    itinerary = data.get('itinerary')
    target_day = data.get('targetDay')
    try:
        target_day = int(target_day) if target_day is not None else None
    except (TypeError, ValueError):
        return jsonify({'error': '指定天數格式不正確。'}), 400
    if not city:
        return jsonify({'error': '請先輸入目的地。'}), 400
    if not isinstance(itinerary, list) or target_day is None:
        return jsonify({'error': '請選擇要調整的行程日期。'}), 400

    updated_itinerary = [dict(day_plan) for day_plan in itinerary]
    target_plan = next((day_plan for day_plan in updated_itinerary if day_plan.get('day') == target_day), None)
    if target_plan is None:
        return jsonify({'error': '找不到指定的行程日期。'}), 400
    date_text = str(target_plan.get('date') or datetime.now().strftime('%Y-%m-%d'))
    try:
        result = generate_gemini_json(f"""
請重新安排 {city} 第 {target_day} 天的旅程，原因是「{reason}」。
原本行程：{json.dumps(target_plan, ensure_ascii=False)}
只替換這一天，請提供至少 4 個、最多 6 個按照時間排序的明確地點或店名。
回傳 JSON：{{"day":{target_day},"date":"{date_text}","title":"標題","description":"調整原因與建議",
"activities":[{{"time":"09:00-11:00","name":"明確地點或店名","place":"可搜尋地點","address":"完整地址","category":"景點","description":"說明"}}]}}
category 只能使用：食物、活動、購物、景點、酒店、交通、其他。
""")
        replanned = normalize_day(result, city, target_day, date_text, include_maps=False)
    except RuntimeError:
        app.logger.error("GEMINI_API_KEY is not configured for replanning.")
        fallback_activity = '室內景點' if reason == '下雨' else '彈性景點'
        replanned = normalize_day({
            'title': f'{city} 第 {target_day} 天（已依{reason}調整）',
            'description': f'因{reason}改安排室內或彈性景點。',
            'activities': [
                {'time': '10:00-11:30', 'name': f'{city} 車站周邊', 'place': f'{city} 車站周邊',
                 'category': '交通', 'description': '保留交通彈性。'},
                {'time': '12:00-13:30', 'name': fallback_activity, 'place': fallback_activity,
                 'category': '活動', 'description': '依突發狀況調整。'},
                {'time': '14:00-15:30', 'name': f'{city} 室內文化景點', 'place': f'{city} 室內文化景點',
                 'category': '景點', 'description': '安排室內行程，降低天候影響。'},
                {'time': '16:00-17:30', 'name': f'{city} 在地商圈', 'place': f'{city} 在地商圈',
                 'category': '購物', 'description': '保留自由探索與購物時間。'},
            ],
        }, city, target_day, date_text, include_maps=False)
    except Exception:
        app.logger.exception("Gemini itinerary replanning failed.")
        return jsonify({'error': 'AI 重新安排失敗，請稍後再試。'}), 502
    updated_itinerary = [
        replanned if day_plan.get('day') == target_day else normalize_day(
            day_plan, city, int(day_plan.get('day', index + 1)), str(day_plan.get('date') or ''),
            include_maps=False
        )
        for index, day_plan in enumerate(updated_itinerary)
    ]
    return jsonify({'itinerary': updated_itinerary, 'message': f'已只調整第 {target_day} 天，其餘行程保持不變。'})


@app.post('/api/recognize_image')
def recognize_image():
    if 'image' not in request.files or not request.files['image'].filename:
        return jsonify({'error': '請選擇圖片。'}), 400
    image = request.files['image']
    image_bytes = image.read()
    if not image_bytes:
        return jsonify({'error': '圖片內容不可為空。'}), 400
    if not image.mimetype or not image.mimetype.startswith('image/'):
        return jsonify({'error': '只支援圖片檔案。'}), 400

    prompt = """
請辨識這張旅遊圖片中的地標、景點或場景。
若無法確定，請誠實說明不確定，不要捏造精確地點。
請推薦 2 個相似或適合延伸探索的景點，每個都要提供可在 Google Maps 搜尋的地點名稱。
回傳 JSON：{"landmarkName":"辨識名稱或未知","description":"辨識說明","similarSpots":[{"name":"景點名稱","emoji":"圖示","reason":"推薦原因","address":"完整地址"}]}
"""
    try:
        result = generate_gemini_json(
            prompt,
            contents=[
                prompt,
                genai.types.Part.from_bytes(data=image_bytes, mime_type=image.mimetype),
            ],
        )
        if not isinstance(result, dict):
            raise ValueError("Gemini returned an invalid image result.")
        result.setdefault('similarSpots', [])
        landmark_name = str(result.get('landmarkName') or '').strip()
        if landmark_name and landmark_name != '未知':
            result['mapsUrl'] = google_maps_search_url(landmark_name)
        for spot in result['similarSpots']:
            if isinstance(spot, dict):
                spot['mapsUrl'] = google_maps_search_url(
                    str(spot.get('address') or spot.get('name') or '').strip()
                )
        return jsonify({**result, 'source': 'gemini'})
    except RuntimeError:
        app.logger.error("GEMINI_API_KEY is not configured.")
        return jsonify({'error': 'Gemini API 尚未設定，請確認 GEMINI_API_KEY 環境變數。'}), 503
    except Exception:
        app.logger.exception("Gemini image recognition failed.")
        return jsonify({'error': 'AI 圖片辨識失敗，請稍後再試。'}), 502

# ==========================================
# 註冊功能 - 支援個人資料
# ==========================================
@app.route('/register', methods=['GET', 'POST'])
def register():
    if request.method == 'POST':
        # 1. 接收所有表單資料
        email = request.form.get('email')
        password = request.form.get('password')
        full_name = (request.form.get('name') or '').strip()
        gender = (request.form.get('gender') or '').strip()
        phone = (request.form.get('phone') or '').strip()

        if not email or not password or not full_name or not phone or gender not in {'male', 'female', 'other'}:
            flash('請完整填寫註冊資料，並選擇有效的性別。', 'warning')
            return redirect(url_for('register'))

        try:
            # 2. 呼叫 Supabase 的註冊 API
            # 將資料打包進 options["data"] 中
            res = supabase.auth.sign_up({
                "email": email,
                "password": password,
                "options": {
                    "data": {
                        "full_name": full_name,
                        "gender": gender,
                        "phone": phone,
                    }
                }
            })

            # 顯示成功訊息並重定向到登入頁面
            flash('註冊信已發送！請前往您的信箱點擊驗證連結。', 'info')
            return redirect(url_for('login'))

        except Exception as e:
            # 處理註冊失敗的情況
            flash(f'註冊失敗: {e}', 'danger')
            return redirect(url_for('register'))

    # 對於 GET 請求，渲染註冊頁面
    return render_template('register.html')

# ==========================================
# 登入功能
# ==========================================
@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        # 獲取表單中的信箱和密碼
        email = request.form['email']
        password = request.form['password']

        try:
            # 使用 Supabase 進行密碼登入
            res = supabase.auth.sign_in_with_password({
                "email": email,
                "password": password
            })
            # 將 Supabase 使用者與 session token 存入 Flask session，供 RLS API 使用
            session['user'] = res.user.email
            session['user_id'] = res.user.id
            session['access_token'] = res.session.access_token
            session['refresh_token'] = res.session.refresh_token

            # [新增] 登入成功時，把 user_metadata 裡面的姓名與性別存入 Session
            user_meta = res.user.user_metadata or {}
            session['name'] = user_meta.get('full_name', res.user.email) # 若無姓名則顯示 email 作為備用
            session['gender'] = user_meta.get('gender')

            # 顯示歡迎訊息並重定向到首頁
            flash(f'歡迎回來, {session["name"]}!', 'success')
            return redirect(url_for('home'))
        except Exception as e:
            # 處理登入失敗的情況
            error_msg = str(e)
            if "Email not confirmed" in error_msg:
                flash('登入失敗：您的帳號尚未開通！請先前往信箱點擊驗證連結。', 'warning')
            else:
                flash('登入失敗：請檢查信箱或密碼是否正確。', 'danger')

            return redirect(url_for('login'))

    # 對於 GET 請求，渲染登入頁面
    return render_template('login.html')

@app.route('/logout')
def logout():
    try:
        # 嘗試從 Supabase 登出
        supabase.auth.sign_out()
    except Exception as e:
        # 記錄登出錯誤，但不影響用戶體驗
        print(f"Supabase 登出錯誤: {e}")

    # 清除 session 中的用戶資訊
    session.pop('user', None)
    session.pop('user_id', None)
    session.pop('access_token', None)
    session.pop('refresh_token', None)
    session.pop('name', None)    # [新增] 清除姓名
    session.pop('gender', None)  # [新增] 清除性別
    # 顯示登出訊息並重定向到首頁
    flash('您已成功登出', 'info')
    return redirect(url_for('home'))


# ==========================================
# 啟動區塊
# ==========================================
if __name__ == "__main__":
    # Render 會提供 PORT 環境變數，若無則預設 5002
    port = int(os.environ.get("PORT", 5002))
    # 生產環境中 debug 應設為 False，但在 Render 介面可透過環境變數控制
    app.run(host='0.0.0.0', port=port)